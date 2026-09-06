import pandas as pd
import numpy as np
from numbers_parser import Document
import re

UP = "/root/.claude/uploads/0753a3d8-6a30-5704-967b-1a152ed9e391/"
PIPE = "/home/claude/work/lisbon_housing_baseline/data/"
OUT = "/home/claude/work/lisbon_housing_baseline/data/clean/"

MONTHS_PT = None

# ---------- helper: month name -> number ----------
MONTH_MAP = {
    'January':1,'February':2,'March':3,'April':4,'May':5,'June':6,
    'July':7,'August':8,'September':9,'October':10,'November':11,'December':12
}
QUARTER_ORDINAL = {'1st':1,'2nd':2,'3rd':3,'4th':4}

def month_label_to_date(label):
    # "December 2022" -> 2022-12-01
    m, y = label.split()
    return pd.Timestamp(year=int(y), month=MONTH_MAP[m], day=1)

def quarter_label_to_date(label):
    # "3rd Quarter 2021" -> 2021-07-01
    parts = label.split()
    ordinal = parts[0]
    year = int(parts[-1])
    q = QUARTER_ORDINAL[ordinal]
    month = (q-1)*3 + 1
    return pd.Timestamp(year=year, month=month, day=1)

# =========================================================
# 1. UNEMPLOYMENT RATE (NUTS I) — quarterly
# =========================================================
doc = Document(UP + "d7048e95-Unemployment_rate_Series_2021___by_Place_of_residence_NUTS__2024_and_Sex_Quarterly.numbers")
table = doc.sheets[0].tables[0]
rows = table.rows(values_only=True)

quarter_row = rows[8]   # e.g. ['2nd Quarter 2026', None, None, '1st Quarter 2026', ...]
sex_row = rows[10]      # ['MF','M','F','MF','M','F', ...]

# build column -> (date, sex) map, skipping col 0 (label) and trailing Nones
col_meta = {}
current_q_label = None
for c in range(1, len(quarter_row)):
    if quarter_row[c] not in (None, ''):
        current_q_label = quarter_row[c]
    sex = sex_row[c]
    if current_q_label and sex in ('MF','M','F'):
        col_meta[c] = (quarter_label_to_date(current_q_label), sex)

region_rows = {
    'Portugal': rows[12],
    'Continente': rows[13],
    'Regiao Autonoma dos Acores': rows[14],
    'Regiao Autonoma da Madeira': rows[15],
}

def parse_val(v):
    if v is None:
        return np.nan
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if 'x' in s.lower():
        return np.nan
    s = s.replace(',', '.')
    s = re.sub(r'[^0-9.\-]', '', s)
    if s in ('', '-', '.'):
        return np.nan
    try:
        return float(s)
    except ValueError:
        return np.nan

records = []
for region, row in region_rows.items():
    for c, (date, sex) in col_meta.items():
        val = parse_val(row[c]) if c < len(row) else np.nan
        records.append({'date': date, 'region': region, 'sex': sex, 'unemployment_rate_pct': val})

unemp = pd.DataFrame(records).sort_values(['region','date','sex']).reset_index(drop=True)
unemp['year'] = unemp['date'].dt.year
unemp['quarter'] = unemp['date'].dt.quarter
unemp.to_csv(OUT + "unemployment_rate_nuts1_quarterly.csv", index=False)
print("unemployment:", unemp.shape, unemp['date'].min(), unemp['date'].max(), unemp['region'].unique())

# =========================================================
# 2. POPULATION (all geographic levels) — annual
# =========================================================
import openpyxl
wb = openpyxl.load_workbook(UP + "cf92f6c3-Populac_a_o_residente_total.xlsx", data_only=True)
ws = wb.worksheets[0]
year_row = 12
years = []
for c in range(3, ws.max_column+1):
    v = ws.cell(row=year_row, column=c).value
    if isinstance(v, int):
        years.append((c, v))
    elif v is None and years:
        break

pop_records = []
for r in range(13, ws.max_row+1):
    geo_level = ws.cell(row=r, column=1).value
    region_name = ws.cell(row=r, column=2).value
    if geo_level is None or region_name is None:
        continue
    if geo_level not in ('NUTS 2024','NUTS I','NUTS II','NUTS III','Município'):
        continue
    for c, yr in years:
        val = ws.cell(row=r, column=c).value
        if val is not None:
            pop_records.append({'geo_level': geo_level, 'region': region_name, 'year': yr, 'population': val})

pop = pd.DataFrame(pop_records)
pop.to_csv(OUT + "population_by_region_annual.csv", index=False)
print("population:", pop.shape, pop['year'].min(), pop['year'].max(), pop['geo_level'].unique())

# =========================================================
# 3. MEDIAN DWELLING SALE PRICE (cities >100k) — quarterly, by typology
# =========================================================
doc2 = Document(UP + "129dbd9b-Median_value_of_dwellings_sales_in_the_last_12_months_Methodology_2018___m__by_Geographic_localization_Cities_with_more_than_100_000_inhabitants_and_Typology_Quarterly.numbers")
table2 = doc2.sheets[0].tables[0]
rows2 = table2.rows(values_only=True)

quarter_row2 = rows2[8]
typology_row2 = rows2[10]

col_meta2 = {}
current_q2 = None
for c in range(1, len(quarter_row2)):
    if quarter_row2[c] not in (None, ''):
        current_q2 = quarter_row2[c]
    typ = typology_row2[c] if c < len(typology_row2) else None
    if current_q2 and typ:
        col_meta2[c] = (quarter_label_to_date(current_q2), typ)

city_rows = {}
for r in range(12, 20):
    label = rows2[r][0]
    if label:
        city_rows[label] = rows2[r]

dwelling_records = []
for city_label, row in city_rows.items():
    # split "1700068: Lisboa" -> code, name
    if ':' in city_label:
        code, name = city_label.split(':', 1)
        name = name.strip()
        code = code.strip()
    else:
        code, name = None, city_label.strip()
    for c, (date, typ) in col_meta2.items():
        val = row[c] if c < len(row) else None
        val = float(val) if isinstance(val, (int, float)) else np.nan
        dwelling_records.append({'date': date, 'city_code': code, 'city': name, 'typology': typ, 'median_price_eur_per_sqm': val})

dwelling = pd.DataFrame(dwelling_records).sort_values(['city','date']).reset_index(drop=True)
dwelling['year'] = dwelling['date'].dt.year
dwelling['quarter'] = dwelling['date'].dt.quarter
dwelling.to_csv(OUT + "median_dwelling_price_city_quarterly.csv", index=False)
print("dwelling price:", dwelling.shape, dwelling['date'].min(), dwelling['date'].max(), dwelling['city'].unique())

# =========================================================
# 4. CONSTRUCTION COST INDEX — monthly, national, by production factor
# =========================================================
doc3 = Document(UP + "3a1ef5b9-Construction_costs_20022022.numbers")
table3 = doc3.sheets[0].tables[0]
rows3 = table3.rows(values_only=True)

cc_records = []
current_month = None
for row in rows3[10:]:
    label, factor, val = row[0], row[1], row[2]
    if label:
        # stop once we hit the metadata footer section
        if label in ('Regularity','Source','First available period','Last available period') or 'Note' in str(label):
            break
        current_month = label
    if factor in ('Total','Materials','Manpower') and current_month:
        v = float(val) if isinstance(val, (int, float)) else np.nan
        cc_records.append({'month_label': current_month, 'production_factor': factor, 'yoy_growth_rate_pct': v})

cc = pd.DataFrame(cc_records)
cc['date'] = cc['month_label'].apply(month_label_to_date)
cc = cc[['date','production_factor','yoy_growth_rate_pct']].sort_values(['production_factor','date']).reset_index(drop=True)
cc.to_csv(OUT + "construction_cost_index_monthly.csv", index=False)
print("construction cost:", cc.shape, cc['date'].min(), cc['date'].max(), cc['production_factor'].unique())
